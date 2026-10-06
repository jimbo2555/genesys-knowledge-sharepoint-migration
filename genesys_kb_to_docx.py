#!/usr/bin/env python3
"""Convert a Genesys Knowledge v2 JSON export into SharePoint-friendly DOCX files."""
from __future__ import annotations
import argparse, csv, io, json, re, subprocess, tempfile
from pathlib import Path
from html import unescape
from urllib.parse import urlparse

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

INVALID = re.compile(r'[\\/:*?"<>|]+')

def safe_name(s, fallback='document'):
    s = INVALID.sub('_', (s or fallback)).strip().strip('.')
    return (s[:150] or fallback)

def fetch_image(url, timeout=30):
    """Download with curl, validate with Pillow, return Word-compatible bytes."""
    suffix = Path(urlparse(url).path).suffix or '.img'
    with tempfile.TemporaryDirectory(prefix='genesys_kb_') as td:
        raw = Path(td) / ('download' + suffix)
        cmd = [
            'curl', '-L', '--fail', '--silent', '--show-error',
            '--connect-timeout', '10', '--max-time', str(timeout),
            '-A', 'Mozilla/5.0 Genesys-KB-Migrator/2.0',
            '-o', str(raw), url
        ]
        p = subprocess.run(cmd, capture_output=True, text=True)
        if p.returncode != 0:
            raise RuntimeError(p.stderr.strip() or f'curl exited {p.returncode}')
        if not raw.exists() or raw.stat().st_size == 0:
            raise RuntimeError('downloaded image is empty')
        data = raw.read_bytes()
        im = Image.open(io.BytesIO(data)); im.load()
        size = im.size
        out = io.BytesIO()
        if im.mode == 'RGBA':
            im.save(out, 'PNG')
        else:
            if im.mode != 'RGB': im = im.convert('RGB')
            im.save(out, 'JPEG', quality=92)
        out.seek(0)
        return out, size

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'), rid)
    run = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rPr.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); rPr.append(underline)
    run.append(rPr); t = OxmlElement('w:t'); t.text = text; run.append(t); hyperlink.append(run)
    paragraph._p.append(hyperlink)

def apply_run(run, node):
    marks = set(node.get('marks') or [])
    run.bold = 'Bold' in marks; run.italic = 'Italic' in marks; run.underline = 'Underline' in marks
    props = node.get('properties') or {}
    sizes = {'Small':9, 'Medium':11, 'Large':14, 'XLarge':18}
    if props.get('fontSize') in sizes: run.font.size = Pt(sizes[props['fontSize']])
    col = props.get('textColor')
    if isinstance(col, str) and re.fullmatch(r'#[0-9A-Fa-f]{6}', col): run.font.color.rgb = RGBColor.from_string(col[1:])

def add_inline(paragraph, blocks, report, source_title):
    for b in blocks or []:
        typ = b.get('type')
        if typ == 'Text':
            n = b.get('text') or {}; text = unescape(n.get('text',''))
            if n.get('hyperlink'): add_hyperlink(paragraph, text, n['hyperlink'])
            else: apply_run(paragraph.add_run(text), n)
        elif typ == 'Image': add_image(paragraph, (b.get('image') or {}).get('url'), report, source_title)
        elif typ == 'Video':
            url = (b.get('video') or {}).get('url',''); paragraph.add_run('Video: '); add_hyperlink(paragraph, url, url)
        else: report.append((source_title, 'warning', f'Unsupported inline block: {typ}'))

def add_image(paragraph, url, report, source_title):
    if not url: return
    try:
        bio, (w, h) = fetch_image(url)
        maxw, maxh = 6.2, 7.5
        ratio = w / h if h else 1
        width = min(maxw, maxh * ratio)
        paragraph.add_run().add_picture(bio, width=Inches(width))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        report.append((source_title, 'image_embedded', url))
    except Exception as e:
        paragraph.add_run('[Image unavailable during migration] ')
        add_hyperlink(paragraph, url, url)
        report.append((source_title, 'image_failed', f'{url} :: {e}'))

def para_style_from_props(p, props):
    ft = (props or {}).get('fontType')
    mapping = {'Heading1':'Heading 1','Heading2':'Heading 2','Heading3':'Heading 3','Heading4':'Heading 4'}
    if ft in mapping: p.style = mapping[ft]
    align = (props or {}).get('align')
    if align == 'Center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'Right': p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def render_block(doc, block, report, title):
    typ = block.get('type')
    if typ == 'Paragraph':
        n = block.get('paragraph') or {}; p = doc.add_paragraph(); para_style_from_props(p, n.get('properties')); add_inline(p, n.get('blocks'), report, title)
    elif typ in ('UnorderedList','OrderedList'):
        n = block.get('list') or {}; style = 'List Bullet' if typ == 'UnorderedList' else 'List Number'
        for item in n.get('blocks') or []:
            if item.get('type') == 'ListItem':
                p = doc.add_paragraph(style=style); add_inline(p, item.get('blocks'), report, title)
                for child in item.get('blocks') or []:
                    if child.get('type') in ('UnorderedList','OrderedList'): render_block(doc, child, report, title)
    elif typ == 'Table':
        rows = (block.get('table') or {}).get('rows') or []
        if not rows: return
        cols = max(len(r.get('cells') or []) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols); table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row.get('cells') or []):
                tc = table.cell(ri,ci); tc.text = ''; p = tc.paragraphs[0]; add_inline(p, cell.get('blocks'), report, title)
                if (cell.get('properties') or {}).get('cellType') == 'HeaderCell':
                    for run in p.runs: run.bold = True
    elif typ == 'Image':
        p = doc.add_paragraph(); add_image(p, (block.get('image') or {}).get('url'), report, title)
    elif typ == 'Video':
        p = doc.add_paragraph(); url = (block.get('video') or {}).get('url',''); p.add_run('Video: '); add_hyperlink(p, url, url)
    else: report.append((title, 'warning', f'Unsupported top-level block: {typ}'))

def choose_record(d, include_drafts=False):
    if d.get('published'): return d['published'], 'published'
    if include_drafts and d.get('draft'): return d['draft'], 'draft'
    return None, None

def convert(input_path, output_dir, include_drafts=False):
    data = json.loads(Path(input_path).read_text(encoding='utf-8'))
    out = Path(output_dir); out.mkdir(parents=True, exist_ok=True)
    report = []; count = 0
    for d in data.get('documents') or []:
        rec, state = choose_record(d, include_drafts)
        if not rec: continue
        title = rec.get('title') or d.get('externalId') or d.get('id') or 'Untitled'
        variations = rec.get('variations') or []
        cat = (rec.get('category') or {}).get('name') or ''
        labels = ', '.join(x.get('name','') for x in (rec.get('labels') or []) if x.get('name'))
        for vi, var in enumerate(variations or [{}], 1):
            doc = Document(); sec = doc.sections[0]
            sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
            doc.core_properties.title = title
            doc.core_properties.subject = 'Migrated from Genesys Knowledge'
            doc.core_properties.keywords = ', '.join(x for x in [cat, labels] if x)
            doc.add_paragraph(style='Title').add_run(title)
            for block in ((var.get('body') or {}).get('blocks') or []): render_block(doc, block, report, title)
            suffix = f'__{safe_name(var.get("name") or str(vi))}' if len(variations) > 1 else ''
            path = out / f'{safe_name(title)}{suffix}.docx'; doc.save(path); count += 1
            report.append((title, 'docx_created', str(path)))
            report.append((title, 'metadata', f'state={state}; genesys_document_id={d.get("id","")}; external_id={d.get("externalId","")}; category={cat}; labels={labels}; variation={var.get("name","")}'))
    rp = out / 'migration_report.csv'
    with rp.open('w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['document','status','detail']); w.writerows(report)
    return count, rp

def main():
    ap = argparse.ArgumentParser(description='Convert Genesys Knowledge v2 JSON exports to DOCX files for SharePoint migration.')
    ap.add_argument('input_json'); ap.add_argument('-o','--output-dir', default='genesys_docx_export'); ap.add_argument('--include-drafts', action='store_true')
    a = ap.parse_args(); n, r = convert(a.input_json, a.output_dir, a.include_drafts); print(f'Created {n} DOCX file(s). Report: {r}')

if __name__ == '__main__': main()
